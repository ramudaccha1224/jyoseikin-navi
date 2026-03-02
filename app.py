import streamlit as st
import json
import os
import io
import unicodedata
from google.genai import Client, types
from dotenv import load_dotenv

load_dotenv()
# ローカル: .env から取得 / Streamlit Cloud: st.secrets から取得
try:
    api_key = st.secrets.get("GEMINI_API_KEY", os.getenv("GEMINI_API_KEY"))
except Exception:
    api_key = os.getenv("GEMINI_API_KEY")
client = Client(api_key=api_key)


# =============================================================
# データロード
# =============================================================
@st.cache_data
def load_all_knowledge():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    with open(os.path.join(base_dir, "form_structures.json"), "r", encoding="utf-8") as f:
        form_map = json.load(f)
    with open(os.path.join(base_dir, "basic_rules.json"), "r", encoding="utf-8") as f:
        rules_and_cases = json.load(f)
    with open(os.path.join(base_dir, "pdf_chunks.json"), "r", encoding="utf-8") as f:
        pdf_chunks = json.load(f)
    return form_map, rules_and_cases, pdf_chunks


# =============================================================
# 半角換算で文字列を切り詰め（日本語＝2、英数字＝1）
# =============================================================
def truncate_half_width(text: str, max_hw: int = 120) -> str:
    count = 0
    for i, ch in enumerate(text):
        w = unicodedata.east_asian_width(ch)
        count += 2 if w in ("F", "W", "A") else 1
        if count > max_hw:
            return text[:i] + "..."
    return text


# =============================================================
# RAG: バイグラムによる関連チャンク抽出（日本語対応）
# =============================================================
def get_relevant_chunks(query: str, pdf_chunks: list, max_chunks: int = 3) -> str:
    scored = []
    for chunk in pdf_chunks:
        content = chunk.get("content", "")
        source  = chunk.get("source", "")
        score = sum(1 for i in range(len(query) - 1) if query[i:i+2] in content)
        if score > 0:
            scored.append((score, content, source))
    scored.sort(key=lambda x: x[0], reverse=True)
    results = [f"[出典: {src}]\n{cont}" for _, cont, src in scored[:max_chunks]]
    return "\n---\n".join(results)


# =============================================================
# システムプロンプト構築（5タイプ判別ロジック統合）
# =============================================================
def build_system_prompt(selected_grant, selected_form, form_map, rules_and_cases, relevant_chunks):
    form_data = form_map.get(selected_form, {})
    return f"""
あなたは『{selected_grant}』専門の助成金申請伴走アドバイザーです。
プロの社会保険労務士として、ユーザーが申請書を正確に完成できるよう伴走支援してください。

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
【最重要：対話の鉄則】
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

■ 文脈最優先の原則（コンテキスト優先）
  - ユーザーの入力が短い（「わからない」「ない」「その予定はない」等）場合、
    または「その」「それ」「そこ」等の代名詞を含む場合は、
    必ず直前の「会話履歴」を参照して意図を解釈すること。
  - JSONデータ内のキーワードを検索して「どの項目ですか？」と聞き返すことは厳禁。

■ 能動的ヒアリング（逆質問）の原則
  - 「支給額は？」等の制度全般に関する質問には、まず基本情報を即答したうえで、
    正確な計算のために必要な情報をAI側から能動的に一問ずつヒアリングすること。

■ 5タイプ判別と回答スタイル
  ▶ タイプ1【チェック型】→ ルールのみ。事例引用厳禁。
  ▶ タイプ2【自由記述型】→ RAG事例を引用して記入見本を作成。
  ▶ タイプ3【数値・計算型】→ 計算式明示。ヒアリング後に具体的計算結果を提示。
  ▶ タイプ4【日付・期間型】→ 期限警告を最優先。
  ▶ タイプ5【選択・フラグ型】→ 定義の違いを解説し選択基準を提示。

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
【対象様式データ】（様式: {selected_form}）
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
{json.dumps(form_data, ensure_ascii=False, indent=2)}

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
【基本ルール・数値定義（支給要領）】
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
{json.dumps(rules_and_cases, ensure_ascii=False)}

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
【活用事例（RAGデータ）— タイプ2【自由記述型】に優先活用】
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
{relevant_chunks if relevant_chunks else "（関連する事例データなし）"}
"""


# =============================================================
# 添削用システムプロンプト構築
# =============================================================
def build_review_prompt(selected_form, form_map, rules_and_cases):
    form_items = form_map.get(selected_form, {}).get("items", [])
    return f"""
あなたは助成金申請書類の専門添削員（プロの社会保険労務士）です。
アップロードされた書類を【様式基準】と【ルール基準】に照らして厳密に添削してください。

【添削手順】
STEP1: 書類の各項目を識別し、【様式基準】のitem_idと照合する。
STEP2: 各記載内容が様式基準の instruction に沿っているか確認する。
STEP3: 数値・日付・計算値が【ルール基準】と矛盾していないか確認する。
STEP4: 結果を ⚠️要修正 / 💡改善提案 / ✅問題なし の3段階で報告。

【様式基準】（{selected_form}）
{json.dumps(form_items, ensure_ascii=False, indent=2)}

【ルール基準】（支給要領）
{json.dumps(rules_and_cases, ensure_ascii=False)}

添削レポートは日本語で、項目ごとに箇条書きでまとめてください。
"""


# =============================================================
# ファイル添削処理（PDF / DOCX / XLSX）
# =============================================================
def review_document(uploaded_file, selected_form, form_map, rules_and_cases):
    file_name  = uploaded_file.name.lower()
    review_sys = build_review_prompt(selected_form, form_map, rules_and_cases)

    if file_name.endswith(".pdf"):
        pdf_bytes = uploaded_file.read()
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[types.Content(role="user", parts=[
                types.Part(inline_data=types.Blob(mime_type="application/pdf", data=pdf_bytes)),
                types.Part(text="このPDF申請書類を添削してください。"),
            ])],
            config=types.GenerateContentConfig(system_instruction=review_sys),
        )
        return response.text

    elif file_name.endswith(".docx"):
        try:
            from docx import Document
            doc  = Document(io.BytesIO(uploaded_file.read()))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            return "❌ `pip install python-docx` が必要です。"
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=f"以下のWord文書を添削してください：\n\n{text}",
            config=types.GenerateContentConfig(system_instruction=review_sys),
        )
        return response.text

    elif file_name.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(uploaded_file.read()))
            all_text = []
            for sn in wb.sheetnames:
                ws = wb[sn]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rs = [str(c) if c is not None else "" for c in row]
                    if any(s.strip() for s in rs):
                        rows.append(" | ".join(rs))
                all_text.append(f"【シート: {sn}】\n" + "\n".join(rows))
            excel_text = "\n\n".join(all_text)
        except ImportError:
            return "❌ `pip install openpyxl` が必要です。"
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=f"以下のExcelシートを添削してください：\n\n{excel_text}",
            config=types.GenerateContentConfig(system_instruction=review_sys),
        )
        return response.text

    return "❌ 対応形式は PDF / Word(.docx) / Excel(.xlsx) のみです。"


# =============================================================
# Gemini 用コンテンツ履歴の構築
# =============================================================
def build_gemini_contents(messages: list, current_prompt: str) -> list:
    contents = []
    for m in messages[:-1]:
        role = "user" if m["role"] == "user" else "model"
        contents.append(types.Content(role=role, parts=[types.Part(text=m["content"])]))
    contents.append(types.Content(role="user", parts=[types.Part(text=current_prompt)]))
    return contents


# =============================================================
# AI応答処理（共通関数化）
# =============================================================
def send_and_stream(prompt: str):
    """ユーザーの質問を処理してストリーミング応答を返す共通関数"""
    relevant_chunks = get_relevant_chunks(prompt, pdf_chunks)
    system_prompt = build_system_prompt(
        st.session_state.selected_grant,
        st.session_state.selected_form,
        form_map, rules_and_cases, relevant_chunks,
    )
    gemini_contents = build_gemini_contents(st.session_state.messages, prompt)

    with st.chat_message("assistant"):
        placeholder = st.empty()
        full = ""
        try:
            for chunk in client.models.generate_content_stream(
                model="gemini-2.5-flash",
                contents=gemini_contents,
                config=types.GenerateContentConfig(system_instruction=system_prompt),
            ):
                if chunk.text:
                    full += chunk.text
                    placeholder.markdown(full + "▌")
            placeholder.markdown(full)
            st.session_state.messages.append({"role": "assistant", "content": full})
        except Exception as e:
            st.error(f"⚠️ エラーが発生しました: {e}")


# =============================================================
# 様式PDFプレビュー（モーダル表示）
# =============================================================
def get_template_path(form_key: str) -> str | None:
    """form_structuresのキーに対応するテンプレートPDFのパスを返す"""
    base_dir = os.path.dirname(os.path.abspath(__file__))
    pdf_path = os.path.join(base_dir, "templates", form_key)
    return pdf_path if os.path.isfile(pdf_path) else None


@st.dialog("確認")
def confirm_reset_dialog():
    """最初の画面に戻る前の確認ダイアログ"""
    st.warning("現在表示されている内容はすべて消去されます。最初の画面に戻りますか？")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("はい", use_container_width=True, type="primary"):
            st.session_state.app_state     = "setup"
            st.session_state.messages      = []
            st.session_state.review_result = ""
            st.session_state.pending_item  = None
            st.rerun()
    with c2:
        if st.button("いいえ", use_container_width=True):
            st.rerun()


@st.dialog("様式プレビュー", width="large")
def show_template_dialog(pdf_path: str):
    """PDFをページごとに画像変換してモーダル表示"""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        st.error("❌ `pip install pymupdf` が必要です。")
        return
    doc = fitz.open(pdf_path)
    for page_num in range(len(doc)):
        page = doc[page_num]
        pix = page.get_pixmap(dpi=150)
        st.image(pix.tobytes("png"), caption=f"ページ {page_num + 1}", use_container_width=True)
    doc.close()


# =============================================================
# Streamlit ページ設定
# =============================================================
st.set_page_config(
    page_title="書類作成エージェント",
    layout="wide",
    page_icon="🛡️",
)

form_map, rules_and_cases, pdf_chunks = load_all_knowledge()

# ── セッション初期化 ──────────────────────────────────────────
_defaults = {
    "app_state":      "setup",
    "messages":       [],
    "selected_grant": "",
    "selected_form":  "",
    "review_result":  "",
    "pending_item":   None,
    "input_key":      0,
}
for k, v in _defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


# =============================================================
# 初期設定画面
# =============================================================
if st.session_state.app_state == "setup":

    st.markdown(
        "<h1 style='text-align:center;'>🛡️ 書類作成エージェント</h1>",
        unsafe_allow_html=True,
    )
    st.markdown(
        "<p style='text-align:center;color:gray;'>"
        "AIと対話しながら、迷わず・正確に助成金申請を完結させます"
        "</p>",
        unsafe_allow_html=True,
    )
    st.divider()

    st.subheader("1. 助成金制度を選択")
    st.session_state.selected_grant = st.selectbox(
        "助成金制度",
        ["人材確保等支援助成金（雇用管理制度・雇用環境整備助成コース）"],
        label_visibility="collapsed",
    )

    st.subheader("2. 相談・添削したい様式を選択")
    form_options = ["全般（様式を特定しない）"] + list(form_map.keys())
    st.session_state.selected_form = st.selectbox(
        "様式", form_options, label_visibility="collapsed",
    )
    st.info("💡 様式を特定するとAIの回答精度と添削の正確さが向上します。", icon="ℹ️")

    if st.button("相談を開始する →", use_container_width=True, type="primary"):
        st.session_state.app_state     = "chat"
        st.session_state.messages      = []
        st.session_state.review_result = ""
        st.rerun()


# =============================================================
# チャット画面 & 添削画面
# =============================================================
else:

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # 左サイドバー（新規チャット・添削モード・様式表示）
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    with st.sidebar:
        st.markdown("### 🛡️ 書類作成エージェント")
        st.divider()

        # ── 添削モード（黄色背景） ──
        st.markdown("""
        <style>
            [data-testid="stSidebar"] [data-testid="stExpander"]:has(summary:first-child) {
                background-color: #FFF3CD;
                border-radius: 8px;
                padding: 2px;
            }
        </style>
        """, unsafe_allow_html=True)

        with st.expander("📝 添削モード"):
            st.caption("申請書類をアップロードして添削します。")
            uploaded_file = st.file_uploader(
                "申請書類", type=["pdf", "docx", "xlsx"], label_visibility="collapsed",
            )
            if uploaded_file:
                st.success(f"📎 {uploaded_file.name}")
                if st.button("🔍 添削実行", type="primary", use_container_width=True):
                    with st.spinner("添削中..."):
                        st.session_state.review_result = review_document(
                            uploaded_file, st.session_state.selected_form,
                            form_map, rules_and_cases,
                        )
                    st.rerun()

        st.divider()

        # ── 最初の画面に戻る（確認ダイアログ付き） ──
        if st.button("← 最初の画面に戻る", use_container_width=True):
            confirm_reset_dialog()

        # ── 様式を画像で表示する ──
        template_path = get_template_path(st.session_state.selected_form)
        if template_path:
            if st.button("📋 様式を画像で表示する", use_container_width=True):
                show_template_dialog(template_path)

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # メインエリア（チャット） + 右カラム（項目一覧）
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

    form_items = form_map.get(st.session_state.selected_form, {}).get("items", [])

    # 右カラムの有無でレイアウトを切り替え
    if form_items:
        col_main, col_right = st.columns([3, 1])
    else:
        col_main = st.container()
        col_right = None

    # ── メインカラム ──────────────────────────────────────────
    with col_main:

        # カスタムCSS（右カラム独立スクロール・ユーザーメッセージ色・様式タイトル）
        st.markdown("""
        <style>
            /* ── 様式タイトル強調 ── */
            .form-title {
                font-size: 22px;
                font-weight: 700;
                color: #FF6B35;
                margin: 0 0 5px 0;
            }

            /* ── 右カラムだけ固定＋独立スクロール ── */
            [data-testid="stColumn"]:has(.right-col-header) > div:first-child {
                position: sticky;
                top: 60px;
                max-height: calc(100vh - 80px);
                overflow-y: auto;
            }

            /* ── ユーザー投稿の背景色（複数セレクタで確実に適用） ── */
            [data-testid="stChatMessage"]:has([data-testid*="user"]),
            [data-testid="stChatMessage"]:has([data-testid*="User"]),
            [data-testid="stChatMessage"][aria-label*="user"] {
                background-color: #d0d0d0 !important;
            }
        </style>
        """, unsafe_allow_html=True)

        # ヘッダー
        st.markdown(f"### 💬 {st.session_state.selected_grant}")
        st.markdown(
            f"<p class='form-title'>📋 {st.session_state.selected_form}</p>",
            unsafe_allow_html=True,
        )

        # 添削レポート（あれば表示）
        if st.session_state.review_result:
            with st.expander("📋 添削レポート", expanded=True):
                st.markdown(st.session_state.review_result)
                if st.button("チャット履歴に追加"):
                    st.session_state.messages.append({
                        "role": "assistant",
                        "content": f"【📋 添削レポート】\n\n{st.session_state.review_result}",
                    })
                    st.session_state.review_result = ""
                    st.rerun()

        st.divider()

        # ── チャット履歴の表示 ────────────────────────────────
        for msg in st.session_state.messages:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])

        # ── 項目ボタンからの自動送信処理 ──────────────────────
        if st.session_state.pending_item is not None:
            item = st.session_state.pending_item
            st.session_state.pending_item = None

            item_id = item.get("item_id", "")
            label   = item.get("label", "")
            prompt  = f"{item_id}「{label}」について教えてください"

            st.session_state.messages.append({"role": "user", "content": prompt})
            with st.chat_message("user"):
                st.markdown(prompt)

            send_and_stream(prompt)
            st.rerun()

        # ── ユーザー入力欄（text_area: 2倍の高さ）────────────
        st.markdown("**質問を入力してください**")
        user_input = st.text_area(
            "入力欄",
            placeholder="例：離職率の計算方法は？ / ③(1)欄には何を書く？",
            height=120,
            label_visibility="collapsed",
            key=f"user_input_{st.session_state.input_key}",
        )

        c1, c2 = st.columns([1, 4])
        with c1:
            submit = st.button("送信", use_container_width=True, type="primary")

        if submit and user_input.strip():
            prompt = user_input.strip()
            st.session_state.messages.append({"role": "user", "content": prompt})
            with st.chat_message("user"):
                st.markdown(prompt)
            send_and_stream(prompt)
            st.session_state.input_key += 1
            st.rerun()

    # ── 右カラム（項目一覧・固定風） ─────────────────────────
    if col_right is not None:
        with col_right:
            st.markdown("""
            <style>
                .right-col-header {
                    font-size: 16px;
                    font-weight: 600;
                    color: #667eea;
                    margin-bottom: 10px;
                }
            </style>
            """, unsafe_allow_html=True)

            st.markdown('<div class="right-col-header">❓ 何について聞きたいですか？</div>', unsafe_allow_html=True)

            for i, item in enumerate(form_items):
                item_id = item.get("item_id", f"項目{i+1}")
                label   = item.get("label", "")
                display = truncate_half_width(f"{item_id}: {label}", 120)
                btn_label = f"📌 {display}"

                if st.button(btn_label, key=f"ri-{i}", use_container_width=True):
                    st.session_state.pending_item = item
                    st.rerun()
